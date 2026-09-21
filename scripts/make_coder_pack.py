"""Build a blind coding pack for a second coder.

    python scripts/make_coder_pack.py --master oshc-master-data-v7-final.xlsx --coder Minhaj

Why a separate file exists at all: the master workbook shows Code 1, Code 2 and
Gap for every record. A second coder who opens it is no longer independent, and
kappa computed afterwards measures nothing. This script emits a workbook that
contains the record IDs and snippets and nothing else, in a shuffled order
under a recorded seed.

Anchor examples are drawn from third-party commenter records, which sit outside
the first-person analysis set, so agreeing on the frame costs no records from
the kappa sample. Where a code has no commenter example, the anchor is taken
from a first-person record and that record's ID is written to the withheld list
so it can be dropped from the kappa calculation.

Owner: Bikram Bhattarai.
"""
from __future__ import annotations

import argparse
import json
import random
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

CODES = {
    "C1": "Purpose of cover",
    "C2": "What is covered",
    "C3": "Accessing healthcare",
    "C4": "Claiming process",
    "C5": "Choosing the right service",
    "C6": "Emergency and urgent help",
    "S": "No comprehension gap: a service complaint, or a disliked policy term the person understood",
}
FILL_ME = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="D9D9D9")
BODY = Font(name="Times New Roman", size=11)
BOLD = Font(name="Times New Roman", size=11, bold=True)


def read_master(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = list(wb["Review Coding"].iter_rows(values_only=True))
    h = [str(x).strip() if x else "" for x in rows[0]]
    idx = {n: h.index(n) for n in h}
    recs = []
    for r in rows[1:]:
        if not r[idx["Review ID"]]:
            continue
        recs.append({
            "id": str(r[idx["Review ID"]]).strip(),
            "snippet": str(r[idx["Snippet"]] or "").strip(),
            "role": str(r[idx["Speaker Role"]] or "").strip(),
            "status": str(r[idx["Analysis Status"]] or "").strip(),
            "code1": str(r[idx["Code 1"]] or "").strip(),
            "gap": str(r[idx["Gap"]] or "").strip(),
        })
    return recs


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--master", required=True)
    ap.add_argument("--coder", required=True)
    ap.add_argument("--seed", type=int, default=20260921)
    ap.add_argument("--out-dir", default=".")
    args = ap.parse_args()

    recs = read_master(Path(args.master))
    included = [r for r in recs if r["status"].startswith("Included")]
    first_person = [r for r in included if r["role"] != "Commenter"]
    commenters = [r for r in included if r["role"] == "Commenter"]

    anchors, withheld = {}, []
    for code in CODES:
        pick = next((r for r in commenters if r["code1"] == code), None)
        if pick is None:
            pick = next((r for r in first_person if r["code1"] == code), None)
            if pick:
                withheld.append(pick["id"])
        if pick:
            anchors[code] = pick

    sample = [r for r in first_person if r["id"] not in withheld]
    rng = random.Random(args.seed)
    order = sample[:]
    rng.shuffle(order)

    wb = openpyxl.Workbook()

    ins = wb.active
    ins.title = "Read me first"
    lines = [
        (f"Coding pack for {args.coder}", True),
        ("", False),
        ("Do not open the master workbook before you finish this file.", True),
        ("This pack deliberately contains no codes. If you see anyone else's codes "
         "before you finish, the agreement statistic we compute afterwards measures "
         "nothing and the work has to be redone.", False),
        ("", False),
        ("What to do", True),
        ("1. Read each snippet on the Coding sheet. Code the text as written. Do not "
         "infer a backstory the words do not state.", False),
        ("2. Enter Code 1, the single best-fitting category. Use the dropdown.", False),
        ("3. Enter Code 2 only if a second category clearly also applies.", False),
        ("4. Enter Gap, Y or N, using the rule below.", False),
        ("5. Send the file back. Do not discuss any record with Bikram until both passes are in.", False),
        ("", False),
        ("Rule 1. Gap: Y or N", True),
        ("Y only when the text shows the person misunderstood something: a question about "
         "how it works, a stated wrong belief ('I thought', 'I didn't know'), or surprise "
         "that reveals a wrong expectation.", False),
        ("N when the person shows they understood what should happen and complains that it "
         "did not happen, happened slowly, or happened badly. Also N when they dislike a "
         "policy term they plainly understand. Complaining about an amount is not by itself "
         "a gap. Asking why the amount was what it was is.", False),
        ("Illustrations, not from the data: 'I assumed the pharmacy would bill my insurer "
         "directly' is Y. 'Third phone call and my claim is still sitting there' is N. "
         "'The dental limit is too low for what it costs' is N.", False),
        ("", False),
        ("Rule 2. Gap = N always takes Code 1 = S. Gap = Y never takes Code 1 = S. "
         "S may appear as Code 2 on a gap record that also contains a service complaint.", True),
        ("", False),
        ("Rule 3. C2 against C4, the boundary most disagreements will sit on", True),
        ("C2 when the confusion is about the policy's terms, the kind that exists before "
         "any claim: what is in or out, limits, caps, waiting periods, what 'covered' means.", False),
        ("C4 when the confusion comes from paying and claiming: paying first, forms, how "
         "much came back, a gap on a particular bill.", False),
        ("When both apply, Code 1 is whichever the person is actually asking or complaining "
         "about, and the other goes in Code 2.", False),
        ("", False),
        ("Rule 4. The other boundaries", True),
        ("C1 against C2: C1 when they do not understand what OSHC is or why they have it. "
         "C2 when they know it is insurance and misjudged its scope.", False),
        ("C3 against C5: C3 when they know what care they need and cannot find, book or "
         "reach it. C5 when they went to the wrong kind of service.", False),
        ("C5 against C6: C6 when the question is how to get emergency or urgent help. "
         "C5 when they used emergency care for something that was not an emergency.", False),
        ("", False),
        ("Categories and agreed examples", True),
    ]
    row = 1
    for text, bold in lines:
        c = ins.cell(row=row, column=1, value=text)
        c.font = BOLD if bold else BODY
        c.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    for code, label in CODES.items():
        ins.cell(row=row, column=1, value=code).font = BOLD
        ins.cell(row=row, column=2, value=label).font = BODY
        a = anchors.get(code)
        ex = ins.cell(row=row, column=3,
                      value=(a["snippet"][:150] if a else "no agreed example, use the rules above"))
        ex.font = Font(name="Times New Roman", size=10, italic=True)
        ex.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    ins.column_dimensions["A"].width = 16
    ins.column_dimensions["B"].width = 34
    ins.column_dimensions["C"].width = 70

    sh = wb.create_sheet("Coding")
    headers = ["Review ID", "Snippet", "Code 1", "Code 2", "Gap"]
    for i, hname in enumerate(headers, start=1):
        c = sh.cell(row=1, column=i, value=hname)
        c.font = BOLD
        c.fill = HEAD
    example = ["EXAMPLE (delete this row)",
               "I thought OSHC covered everything so I never checked the policy.",
               "C2", "C1", "Y"]
    for i, v in enumerate(example, start=1):
        c = sh.cell(row=2, column=i, value=v)
        c.font = Font(name="Times New Roman", size=11, italic=True)
    for j, r in enumerate(order, start=3):
        sh.cell(row=j, column=1, value=r["id"]).font = BODY
        sc = sh.cell(row=j, column=2, value=r["snippet"])
        sc.font = BODY
        sc.alignment = Alignment(wrap_text=True, vertical="top")
        for col in (3, 4, 5):
            cell = sh.cell(row=j, column=col)
            cell.fill = FILL_ME
            cell.font = BODY

    last = len(order) + 2
    dv_code = DataValidation(type="list", formula1='"' + ",".join(CODES) + '"', allow_blank=True)
    dv_gap = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    sh.add_data_validation(dv_code)
    sh.add_data_validation(dv_gap)
    dv_code.add(f"C3:D{last}")
    dv_gap.add(f"E3:E{last}")
    for col, w in zip("ABCDE", (16, 95, 12, 12, 8)):
        sh.column_dimensions[col].width = w
    sh.freeze_panes = "A2"

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"coder_pack_{args.coder.lower()}.xlsx"
    wb.save(out)

    meta = {
        "coder": args.coder,
        "built_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "master": Path(args.master).name,
        "shuffle_seed": args.seed,
        "records_in_pack": len(order),
        "first_person_total": len(first_person),
        "withheld_as_anchors": withheld,
        "codes_without_commenter_anchor": [c for c in CODES if anchors.get(c) in (None,)
                                           or anchors[c]["role"] != "Commenter"],
    }
    (out_dir / f"coder_pack_{args.coder.lower()}_manifest.json").write_text(
        json.dumps(meta, indent=2), encoding="utf-8")

    print(f"pack written      : {out}")
    print(f"records in pack   : {len(order)} of {len(first_person)} first-person records")
    print(f"withheld as anchor: {', '.join(withheld) or 'none'}")
    print(f"shuffle seed      : {args.seed}")
    print()
    print("Send only this file. Do not send the master workbook.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
