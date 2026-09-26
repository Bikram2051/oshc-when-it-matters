"""Convert a Microsoft Forms Excel export into a probe CSV for seminar.score.

The form has three free-text questions whose titles start with a tag:
    [ANSWER]     a question the assistant should answer
    [REFUSE]     a question it should refuse and redirect
    [EMERGENCY]  a message where someone may need emergency care now

Each respondent's answers become up to three probe rows; blank answers are
skipped. Only the tagged answers are kept. Names, emails and timestamps in the
export are never copied, and author is always written as "anonymous".

Run from the repo root:

    python -m seminar.form_to_probes <export.xlsx> seminar/class_probes.csv --split class --prefix C
"""

from __future__ import annotations

import argparse
import csv
import sys
from collections import Counter
from pathlib import Path

TAGS = {"[ANSWER]": "answer", "[REFUSE]": "refuse", "[EMERGENCY]": "escalate"}
LABEL_ORDER = ("answer", "refuse", "escalate")


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Microsoft Forms export to probe CSV.")
    parser.add_argument("xlsx", type=Path, help="the .xlsx downloaded from Microsoft Forms")
    parser.add_argument("out", type=Path, help="probe CSV to write")
    parser.add_argument("--split", required=True, help="split name, for example class or backup")
    parser.add_argument("--prefix", required=True, help="id prefix, for example C or F (B is the dev set)")
    args = parser.parse_args(argv)

    if args.prefix.upper() == "B":
        sys.exit("Prefix B belongs to the dev set. Use C for the class set or F for the backup set.")

    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("openpyxl is not installed. Run: uv pip install openpyxl  (or: python -m pip install openpyxl)")

    sheet = load_workbook(args.xlsx, read_only=True, data_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        sys.exit(f"{args.xlsx}: the first sheet is empty.")

    columns: dict[str, int] = {}
    for index, title in enumerate(header):
        title_text = str(title or "").strip().upper()
        for tag, label in TAGS.items():
            if title_text.startswith(tag):
                if label in columns:
                    sys.exit(f"Two columns start with {tag}. Each tag must appear once.")
                columns[label] = index
    missing = [tag for tag, label in TAGS.items() if label not in columns]
    if missing:
        sys.exit(f"No question title starts with {', '.join(missing)}. Check the titles in the form.")

    probes: list[tuple[str, str]] = []
    respondents = 0
    blanks = 0
    for row in rows:
        if row is None or all(cell is None for cell in row):
            continue
        respondents += 1
        for label in LABEL_ORDER:
            index = columns[label]
            value = row[index] if index < len(row) else None
            text = " ".join(str(value).split()) if value is not None else ""
            if text:
                probes.append((text, label))
            else:
                blanks += 1

    if not probes:
        sys.exit("No answers found in the export.")

    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["id", "text", "expected_action", "author", "split"])
        for number, (text, label) in enumerate(probes, start=1):
            writer.writerow([f"{args.prefix.upper()}{number:03d}", text, label, "anonymous", args.split])

    counts = Counter(label for _, label in probes)
    duplicates = len(probes) - len({text.lower() for text, _ in probes})
    print(f"Respondents: {respondents}   blank answers skipped: {blanks}")
    print("Probes: " + ", ".join(f"{label} {counts[label]}" for label in LABEL_ORDER) + f"   total {len(probes)}")
    if duplicates:
        print(f"Note: {duplicates} duplicate text(s) kept as separate rows.")
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
